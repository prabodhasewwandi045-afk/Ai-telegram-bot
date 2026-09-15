import os
import requests
from docx import Document

from telegram import Update
from telegram.ext import (
    Application,
    CommandHandler,
    MessageHandler,
    ContextTypes,
    ConversationHandler,
    filters,
)

TELEGRAM_TOKEN = os.environ["TELEGRAM_BOT_TOKEN"]
GEMINI_KEY = os.environ["GEMINI_API_KEY"]

GEMINI_URL = (
    "https://generativelanguage.googleapis.com/v1beta/models/"
    "gemini-3.6-flash:generateContent"
)

# CV steps
NAME, PHONE, EMAIL, JOB, EDUCATION, EXPERIENCE, SKILLS, ADDRESS = range(8)


async def start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await update.message.reply_text(
        "🤖 Hello! මම ඔයාගේ AI Bot.\n\n"
        "සාමාන්‍ය ප්‍රශ්න මගෙන් අහන්න.\n"
        "CV එකක් හදන්න /cv කියලා යවන්න."
    )


async def chat(update: Update, context: ContextTypes.DEFAULT_TYPE):
    text = update.message.text

    try:
        response = requests.post(
            GEMINI_URL,
            headers={
                "x-goog-api-key": GEMINI_KEY,
                "Content-Type": "application/json",
            },
            json={
                "contents": [
                    {
                        "parts": [
                            {
                                "text": f"""
ඔබ helpful AI assistant කෙනෙකි.

පරිශීලකයා සිංහල හෝ Singlish වලින් අසනවා නම්
සරල පැහැදිලි සිංහලෙන් පිළිතුරු දෙන්න.

පරිශීලකයා English වලින් අසනවා නම්
English වලින් පිළිතුරු දෙන්න.

ප්‍රශ්නය:
{text}
"""
                            }
                        ]
                    }
                ]
            },
            timeout=60,
        )

        data = response.json()

        if response.ok:
            answer = data["candidates"][0]["content"]["parts"][0]["text"]
            await update.message.reply_text(answer)
        else:
            error = data.get("error", {}).get("message", "Unknown error")
            await update.message.reply_text(f"❌ Gemini Error:\n{error}")

    except Exception as e:
        await update.message.reply_text(f"❌ Error:\n{e}")


async def cv_start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data["cv"] = {}

    await update.message.reply_text(
        "📄 CV Generator\n\n"
        "පළමුව ඔබගේ සම්පූර්ණ නම එවන්න."
    )

    return NAME


async def cv_name(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data["cv"]["name"] = update.message.text
    await update.message.reply_text("📱 Phone number එක එවන්න.")
    return PHONE


async def cv_phone(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data["cv"]["phone"] = update.message.text
    await update.message.reply_text("📧 Email address එක එවන්න.")
    return EMAIL


async def cv_email(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data["cv"]["email"] = update.message.text
    await update.message.reply_text("🎯 Apply කරන job position එක එවන්න.")
    return JOB


async def cv_job(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data["cv"]["job"] = update.message.text
    await update.message.reply_text("🎓 Education details එවන්න.")
    return EDUCATION


async def cv_education(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data["cv"]["education"] = update.message.text
    await update.message.reply_text("💼 Work experience එක එවන්න.")
    return EXPERIENCE


async def cv_experience(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data["cv"]["experience"] = update.message.text
    await update.message.reply_text("🛠️ Main skills ටික එවන්න.")
    return SKILLS


async def cv_skills(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data["cv"]["skills"] = update.message.text
    await update.message.reply_text("📍 Address එක එවන්න.")
    return ADDRESS


async def cv_address(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data["cv"]["address"] = update.message.text

    cv = context.user_data["cv"]

    try:
        document = Document()

        document.add_heading(cv["name"], 0)
        document.add_paragraph(cv["job"])

        document.add_heading("Contact Information", level=1)
        document.add_paragraph(f"Phone: {cv['phone']}")
        document.add_paragraph(f"Email: {cv['email']}")
        document.add_paragraph(f"Address: {cv['address']}")

        document.add_heading("Education", level=1)
        document.add_paragraph(cv["education"])

        document.add_heading("Work Experience", level=1)
        document.add_paragraph(cv["experience"])

        document.add_heading("Skills", level=1)
        document.add_paragraph(cv["skills"])

        filename = "Professional_CV.docx"
        document.save(filename)

        await update.message.reply_text(
            "✅ CV එක සාර්ථකව සෑදුවා!\n"
            "📄 පහළින් CV file එක download කරන්න."
        )

        with open(filename, "rb") as file:
            await update.message.reply_document(
                document=file,
                filename=filename,
            )

    except Exception as e:
        await update.message.reply_text(
            f"❌ CV එක හදද්දී error එකක් ආවා:\n{e}"
        )

    context.user_data.clear()
    return ConversationHandler.END


async def cv_cancel(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data.clear()

    await update.message.reply_text(
        "❌ CV creation cancelled."
    )

    return ConversationHandler.END


cv_handler = ConversationHandler(
    entry_points=[CommandHandler("cv", cv_start)],
    states={
        NAME: [MessageHandler(filters.TEXT & ~filters.COMMAND, cv_name)],
        PHONE: [MessageHandler(filters.TEXT & ~filters.COMMAND, cv_phone)],
        EMAIL: [MessageHandler(filters.TEXT & ~filters.COMMAND, cv_email)],
        JOB: [MessageHandler(filters.TEXT & ~filters.COMMAND, cv_job)],
        EDUCATION: [MessageHandler(filters.TEXT & ~filters.COMMAND, cv_education)],
        EXPERIENCE: [MessageHandler(filters.TEXT & ~filters.COMMAND, cv_experience)],
        SKILLS: [MessageHandler(filters.TEXT & ~filters.COMMAND, cv_skills)],
        ADDRESS: [MessageHandler(filters.TEXT & ~filters.COMMAND, cv_address)],
    },
    fallbacks=[
        CommandHandler("cancel", cv_cancel)
    ],
)


app = Application.builder().token(TELEGRAM_TOKEN).build()

app.add_handler(CommandHandler("start", start))
app.add_handler(cv_handler)
app.add_handler(MessageHandler(filters.TEXT & ~filters.COMMAND, chat))

print("🤖 Telegram Gemini AI Bot is running...")

app.run_polling()
